import sqlite3

from config import DB_PATH, DEFAULT_WORLD_ID


_conn: sqlite3.Connection | None = None


def get_setting(key: str) -> str | None:
    conn = _get_conn()
    row = conn.execute(
        "SELECT value FROM settings WHERE key = ?", (key,)
    ).fetchone()
    return row["value"] if row else None


def set_setting(key: str, value: str):
    conn = _get_conn()
    conn.execute(
        "INSERT INTO settings (key, value) VALUES (?, ?) "
        "ON CONFLICT(key) DO UPDATE SET value = excluded.value",
        (key, value),
    )
    conn.commit()


def get_all_settings() -> dict[str, str]:
    conn = _get_conn()
    rows = conn.execute("SELECT key, value FROM settings").fetchall()
    return {r["key"]: r["value"] for r in rows}


def _get_conn() -> sqlite3.Connection:
    global _conn
    if _conn is None:
        _conn = sqlite3.connect(DB_PATH, check_same_thread=False)
        _conn.row_factory = sqlite3.Row
        _conn.execute("PRAGMA journal_mode=WAL")
        _conn.execute("PRAGMA synchronous=NORMAL")
        _conn.execute(
            """
            CREATE TABLE IF NOT EXISTS worlds (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                description TEXT DEFAULT '',
                created_at TEXT NOT NULL DEFAULT (datetime('now'))
            )
            """
        )
        _conn.execute(
            """
            CREATE TABLE IF NOT EXISTS staging (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                content TEXT NOT NULL,
                source TEXT NOT NULL,
                type TEXT NOT NULL CHECK(type IN ('chat', 'docx', 'text')),
                created_at TEXT NOT NULL DEFAULT (datetime('now')),
                is_processed INTEGER NOT NULL DEFAULT 0
            )
            """
        )
        _conn.execute(
            """
            CREATE TABLE IF NOT EXISTS chat_history (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                query TEXT NOT NULL,
                answer TEXT NOT NULL,
                created_at TEXT NOT NULL DEFAULT (datetime('now'))
            )
            """
        )
        _conn.execute(
            """
            CREATE TABLE IF NOT EXISTS chats (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                title TEXT NOT NULL,
                created_at TEXT NOT NULL DEFAULT (datetime('now')),
                updated_at TEXT NOT NULL DEFAULT (datetime('now'))
            )
            """
        )
        _conn.execute(
            """
            CREATE TABLE IF NOT EXISTS indexing_stats (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                doc_count INTEGER NOT NULL,
                duration_seconds REAL NOT NULL,
                created_at TEXT NOT NULL DEFAULT (datetime('now'))
            )
            """
        )
        _conn.execute(
            """
            CREATE TABLE IF NOT EXISTS settings (
                key TEXT PRIMARY KEY,
                value TEXT NOT NULL
            )
            """
        )
        _migrate()
        _conn.commit()
    return _conn


def _migrate():
    _migrate_worlds_default()
    _migrate_chat_id()
    _migrate_staging_types()
    _migrate_world_id_columns()


def _migrate_worlds_default():
    row = _conn.execute("SELECT COUNT(*) FROM worlds").fetchone()[0]
    if row == 0:
        _conn.execute(
            "INSERT INTO worlds (name, description) VALUES (?, ?)",
            ("Default", "Default world"),
        )


def _migrate_chat_id():
    try:
        _conn.execute(
            "ALTER TABLE chat_history ADD COLUMN chat_id INTEGER REFERENCES chats(id)"
        )
    except sqlite3.OperationalError:
        pass

    row = _conn.execute("SELECT COUNT(*) FROM chats").fetchone()[0]
    if row == 0:
        _conn.execute("INSERT INTO chats (title) VALUES (?)", ("General",))
        default_id = _conn.execute("SELECT last_insert_rowid()").fetchone()[0]
        _conn.execute(
            "UPDATE chat_history SET chat_id = ? WHERE chat_id IS NULL",
            (default_id,),
        )


def _migrate_staging_types():
    try:
        _conn.execute(
            "ALTER TABLE staging ADD COLUMN _migrated INTEGER DEFAULT 0"
        )
    except sqlite3.OperationalError:
        return

    _conn.execute(
        """CREATE TABLE staging_new (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            content TEXT NOT NULL,
            source TEXT NOT NULL,
            type TEXT NOT NULL CHECK(type IN ('chat', 'docx', 'text', 'logseq')),
            created_at TEXT NOT NULL DEFAULT (datetime('now')),
            is_processed INTEGER NOT NULL DEFAULT 0
        )"""
    )
    _conn.execute(
        "INSERT INTO staging_new (id, content, source, type, created_at, is_processed) "
        "SELECT id, content, source, type, created_at, is_processed FROM staging"
    )
    _conn.execute("DROP TABLE staging")
    _conn.execute("ALTER TABLE staging_new RENAME TO staging")


def _migrate_world_id_columns():
    for table, col_type in [
        ("staging", "INTEGER"),
        ("chat_history", "INTEGER"),
        ("chats", "INTEGER"),
        ("indexing_stats", "INTEGER"),
    ]:
        try:
            _conn.execute(
                f"ALTER TABLE {table} ADD COLUMN world_id {col_type} "
                f"REFERENCES worlds(id) DEFAULT {DEFAULT_WORLD_ID}"
            )
        except sqlite3.OperationalError:
            pass

    default_id = _conn.execute(
        "SELECT id FROM worlds ORDER BY id LIMIT 1"
    ).fetchone()[0]

    for table in ["staging", "chat_history", "chats", "indexing_stats"]:
        _conn.execute(
            f"UPDATE {table} SET world_id = ? WHERE world_id IS NULL",
            (default_id,),
        )


def init_db():
    _get_conn()


# --- Worlds API ---


def create_world(name: str, description: str = "") -> int:
    conn = _get_conn()
    cur = conn.execute(
        "INSERT INTO worlds (name, description) VALUES (?, ?)",
        (name, description),
    )
    conn.commit()
    return cur.lastrowid


def list_worlds() -> list[dict]:
    conn = _get_conn()
    rows = conn.execute(
        "SELECT id, name, description, created_at FROM worlds ORDER BY id"
    ).fetchall()
    return [dict(r) for r in rows]


def get_world(world_id: int) -> dict | None:
    conn = _get_conn()
    row = conn.execute(
        "SELECT id, name, description, created_at FROM worlds WHERE id = ?",
        (world_id,),
    ).fetchone()
    return dict(row) if row else None


def rename_world(world_id: int, name: str):
    conn = _get_conn()
    conn.execute("UPDATE worlds SET name = ? WHERE id = ?", (name, world_id))
    conn.commit()


def delete_world(world_id: int):
    conn = _get_conn()
    conn.execute("DELETE FROM staging WHERE world_id = ?", (world_id,))
    conn.execute("DELETE FROM chat_history WHERE world_id = ?", (world_id,))
    conn.execute("DELETE FROM chats WHERE world_id = ?", (world_id,))
    conn.execute("DELETE FROM indexing_stats WHERE world_id = ?", (world_id,))
    conn.execute("DELETE FROM worlds WHERE id = ?", (world_id,))
    conn.commit()


def get_world_stats(world_id: int) -> dict:
    conn = _get_conn()
    staging_total = conn.execute(
        "SELECT COUNT(*) FROM staging WHERE world_id = ?", (world_id,)
    ).fetchone()[0]
    staging_pending = conn.execute(
        "SELECT COUNT(*) FROM staging WHERE world_id = ? AND is_processed = 0",
        (world_id,),
    ).fetchone()[0]
    chat_count = conn.execute(
        "SELECT COUNT(*) FROM chats WHERE world_id = ?", (world_id,)
    ).fetchone()[0]
    return {
        "staging_total": staging_total,
        "staging_pending": staging_pending,
        "chat_count": chat_count,
    }


# --- Staging API ---


def add_to_buffer(content: str, source: str, type_: str, world_id: int = DEFAULT_WORLD_ID):
    conn = _get_conn()
    conn.execute(
        "INSERT INTO staging (content, source, type, world_id) VALUES (?, ?, ?, ?)",
        (content, source, type_, world_id),
    )
    conn.commit()


def bulk_add_to_buffer(records: list[tuple[str, str, str, str | None]], world_id: int = DEFAULT_WORLD_ID):
    conn = _get_conn()
    conn.executemany(
        "INSERT INTO staging (content, source, type, created_at, world_id) VALUES (?, ?, ?, COALESCE(?, datetime('now')), ?)",
        [r + (world_id,) for r in records],
    )
    conn.commit()


def get_unprocessed(world_id: int = DEFAULT_WORLD_ID) -> list[dict]:
    conn = _get_conn()
    rows = conn.execute(
        "SELECT * FROM staging WHERE is_processed = 0 AND world_id = ? ORDER BY id",
        (world_id,),
    ).fetchall()
    return [dict(r) for r in rows]


def mark_processed(ids: list[int]):
    if not ids:
        return
    conn = _get_conn()
    placeholders = ",".join("?" * len(ids))
    conn.execute(
        f"UPDATE staging SET is_processed = 1 WHERE id IN ({placeholders})", ids
    )
    conn.commit()


def get_stats(world_id: int = DEFAULT_WORLD_ID) -> dict:
    conn = _get_conn()
    total = conn.execute(
        "SELECT COUNT(*) FROM staging WHERE world_id = ?", (world_id,)
    ).fetchone()[0]
    pending = conn.execute(
        "SELECT COUNT(*) FROM staging WHERE world_id = ? AND is_processed = 0",
        (world_id,),
    ).fetchone()[0]
    return {"total": total, "pending": pending}


# --- Chats API ---


def create_chat(title: str = "New Chat", world_id: int = DEFAULT_WORLD_ID) -> int:
    conn = _get_conn()
    cur = conn.execute(
        "INSERT INTO chats (title, world_id) VALUES (?, ?)", (title, world_id)
    )
    conn.commit()
    return cur.lastrowid


def rename_chat(chat_id: int, title: str):
    conn = _get_conn()
    conn.execute(
        "UPDATE chats SET title = ?, updated_at = datetime('now') WHERE id = ?",
        (title, chat_id),
    )
    conn.commit()


def delete_chat(chat_id: int):
    conn = _get_conn()
    conn.execute("DELETE FROM chat_history WHERE chat_id = ?", (chat_id,))
    conn.execute("DELETE FROM chats WHERE id = ?", (chat_id,))
    conn.commit()


def list_chats(world_id: int = DEFAULT_WORLD_ID) -> list[dict]:
    conn = _get_conn()
    rows = conn.execute(
        "SELECT id, title, created_at FROM chats WHERE world_id = ? ORDER BY updated_at DESC",
        (world_id,),
    ).fetchall()
    return [dict(r) for r in rows]


def get_chat_title(chat_id: int) -> str | None:
    conn = _get_conn()
    row = conn.execute(
        "SELECT title FROM chats WHERE id = ?", (chat_id,)
    ).fetchone()
    return row["title"] if row else None


def save_chat(query: str, answer: str, chat_id: int | None = None, world_id: int = DEFAULT_WORLD_ID):
    conn = _get_conn()
    conn.execute(
        "INSERT INTO chat_history (query, answer, chat_id, world_id) VALUES (?, ?, ?, ?)",
        (query, answer, chat_id, world_id),
    )
    conn.commit()


def load_chat_history(chat_id: int, limit: int = 50) -> list[dict]:
    conn = _get_conn()
    rows = conn.execute(
        "SELECT * FROM chat_history WHERE chat_id = ? ORDER BY id DESC LIMIT ?",
        (chat_id, limit),
    ).fetchall()
    return [dict(r) for r in reversed(rows)]


def load_recent_chat_context(chat_id: int, n: int = 5) -> str:
    conn = _get_conn()
    rows = conn.execute(
        "SELECT query, answer FROM chat_history WHERE chat_id = ? ORDER BY id DESC LIMIT ?",
        (chat_id, n),
    ).fetchall()
    rows = list(reversed(rows))
    parts: list[str] = []
    for r in rows:
        parts.append(f"User: {r['query']}")
        parts.append(f"Assistant: {r['answer']}")
    return "\n".join(parts)


# --- Indexing stats ---


def save_indexing_run(doc_count: int, duration: float, world_id: int = DEFAULT_WORLD_ID):
    conn = _get_conn()
    conn.execute(
        "INSERT INTO indexing_stats (doc_count, duration_seconds, world_id) VALUES (?, ?, ?)",
        (doc_count, duration, world_id),
    )
    conn.commit()


def get_indexing_estimate(pending_count: int, world_id: int = DEFAULT_WORLD_ID) -> str | None:
    conn = _get_conn()
    row = conn.execute(
        "SELECT SUM(doc_count), SUM(duration_seconds), COUNT(*) FROM indexing_stats WHERE world_id = ?",
        (world_id,),
    ).fetchone()
    if not row or row[2] == 0 or row[0] is None or row[0] == 0:
        return None
    total_docs, total_secs, runs = row
    avg_secs_per_doc = total_secs / total_docs
    est_secs = avg_secs_per_doc * pending_count
    return f"~{est_secs:.0f}s ({runs} previous run{'s' if runs != 1 else ''})"


def parse_docx(filepath: str) -> str:
    from docx import Document

    doc = Document(filepath)
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())